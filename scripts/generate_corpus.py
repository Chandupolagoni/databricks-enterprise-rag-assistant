"""Generate the synthetic document corpus for the Enterprise Knowledge Assistant.

All content is fictional. "Meridian Trust & Co." is a made-up financial services firm
used only to give the RAG pipeline realistic, multi-format, multi-business-unit data.

Outputs:
  data/raw/<business_unit>/<doc_id>.<ext>   (pdf, docx, html, md, txt)
  data/metadata/documents.csv               (doc-level metadata used for RBAC + citations)

Run:  python scripts/generate_corpus.py
"""
from __future__ import annotations

import csv
import html
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw"
META = ROOT / "data" / "metadata" / "documents.csv"

# (doc_id, business_unit, category, access_level, format, title, sections)
# sections: list of (heading, paragraph) tuples
DOCS: list[tuple] = [
    # ------------------------------------------------------------------ compliance
    ("cmp_aml_kyc_policy", "compliance", "financial_policy", "internal", "pdf",
     "Anti-Money Laundering and KYC Policy",
     [("Purpose", "This policy sets the minimum Know Your Customer (KYC) and Anti-Money Laundering (AML) standards for every business unit of Meridian Trust & Co. The policy is owned by the Compliance department and is governed by the Bank Secrecy Act and the USA PATRIOT Act."),
      ("Customer Due Diligence", "Every new client must complete Customer Due Diligence (CDD) before an account is opened. CDD requires verified legal name, date of birth, residential address and a government identification document. Beneficial owners holding 25 percent or more of a legal entity must also be identified."),
      ("Enhanced Due Diligence", "Enhanced Due Diligence (EDD) applies to politically exposed persons, clients in high-risk jurisdictions and any client whose expected annual activity exceeds 10 million USD. EDD reviews must be refreshed every 12 months and approved by a Compliance Officer."),
      ("Suspicious Activity", "Employees who detect unusual activity must raise an internal alert in the Case Manager tool within 24 hours. The Financial Intelligence Unit decides whether to file a Suspicious Activity Report (SAR). A SAR must be filed with FinCEN within 30 calendar days of initial detection. Employees must never tell a client that a SAR is being considered; this is known as tipping off and is prohibited."),
      ("Record Retention", "KYC records must be retained for five years after the account is closed.")]),
    ("cmp_sar_filing_sop", "compliance", "sop", "confidential", "docx",
     "SOP: Suspicious Activity Report Filing",
     [("Scope", "This standard operating procedure applies to analysts in the Financial Intelligence Unit (FIU). It requires the AML and KYC Policy."),
      ("Step 1 - Triage", "Open the alert in Case Manager, confirm it is not a duplicate, and assign a priority. Priority 1 alerts involve possible terrorist financing and must be escalated to the BSA Officer within 4 hours."),
      ("Step 2 - Investigation", "Collect 12 months of transaction history, KYC profile and any adverse media. Document the narrative using the who, what, when, where, why and how structure."),
      ("Step 3 - Decision", "A second analyst reviews the case. If the reviewers disagree, the BSA Officer makes the final decision."),
      ("Step 4 - Filing", "File the SAR through the BSA E-Filing system within 30 calendar days of detection. If no suspect is identified, the deadline may be extended by 30 days, but never beyond 60 days total. Continuing activity reviews are performed every 90 days.")]),
    ("cmp_complaints_manual", "compliance", "complaints_manual", "public", "html",
     "Client Complaints Handling Manual",
     [("Definition", "A complaint is any expression of dissatisfaction from a client, written or verbal, about a product, service or employee of Meridian Trust & Co."),
      ("Acknowledgement", "All complaints must be logged in the Complaints Register and acknowledged to the client within 2 business days."),
      ("Resolution Timeline", "Standard complaints must be resolved within 15 business days. Complaints involving alleged fraud, discrimination or a regulatory breach are classified as escalated and must be routed to Compliance within 1 business day."),
      ("Root Cause", "Each quarter the Complaints Committee reviews root causes and reports trends to the Operational Risk Committee.")]),
    ("cmp_gifts_entertainment", "compliance", "financial_policy", "public", "md",
     "Gifts and Entertainment Policy",
     [("Limits", "Employees may accept gifts with a value of up to 100 USD per giver per calendar year. Business entertainment up to 250 USD per person per event is allowed if the giver attends."),
      ("Pre-approval", "Any gift or entertainment above these limits requires pre-approval from the employee's manager and Compliance through the Gifts Register."),
      ("Government Officials", "Gifts of any value to government officials are prohibited without written approval from the Chief Compliance Officer.")]),
    ("cmp_personal_trading", "compliance", "financial_policy", "internal", "pdf",
     "Personal Trading Policy",
     [("Pre-clearance", "Access persons must pre-clear personal trades in individual securities through the Personal Trading portal. Pre-clearance approval is valid until the end of the next business day."),
      ("Holding Period", "Securities purchased in a personal account must be held for at least 30 calendar days. Short-term profits made within 30 days must be disgorged to charity."),
      ("Restricted List", "Trading in any security on the Restricted List is prohibited. The Restricted List is maintained by Compliance and governed by SEC Rule 204A-1.")]),

    # ------------------------------------------------------------------ operations
    ("ops_wire_transfer_sop", "operations", "sop", "internal", "docx",
     "SOP: Outgoing Wire Transfer Processing",
     [("Scope", "Applies to the Payments Operations team for all outgoing domestic and international wires. This SOP requires the AML and KYC Policy."),
      ("Callback Verification", "Any wire instruction received by email or fax must be verified by a callback to a phone number already on file. Never use a phone number contained in the instruction itself."),
      ("Dual Authorization", "Wires above 250,000 USD require dual authorization: a maker and a checker who are different employees. Wires above 5 million USD also require approval from a Payments Operations manager."),
      ("Cut-off Times", "The cut-off time for same-day domestic Fedwire payments is 4:30 PM Eastern Time. The cut-off for international SWIFT payments is 2:00 PM Eastern Time."),
      ("Sanctions Screening", "All wires are screened against the OFAC sanctions list before release. A potential match must be cleared by the Sanctions team before the payment can proceed.")]),
    ("ops_trade_settlement_kb", "operations", "knowledge_base", "internal", "html",
     "KB: Trade Settlement Failures",
     [("Overview", "US equities settle on T+1. A failed settlement occurs when securities or cash are not delivered on the settlement date."),
      ("Common Causes", "The most common causes are mismatched standing settlement instructions, insufficient securities in the client account and late trade affirmation."),
      ("Resolution", "Contact the counterparty within 2 hours of the fail notification. If the fail persists for 3 days, escalate to the Settlements team lead. Fails older than 5 days must be reported to the Operational Risk team."),
      ("CSDR Penalties", "For European trades, cash penalties under CSDR are applied daily and are charged back to the responsible desk.")]),
    ("ops_corporate_actions_sop", "operations", "sop", "internal", "md",
     "SOP: Voluntary Corporate Actions",
     [("Notification", "The Corporate Actions team must notify affected clients of a voluntary corporate action, such as a tender offer or rights issue, within 1 business day of receiving the announcement."),
      ("Client Elections", "The internal election deadline is 2 business days before the market deadline. Clients who do not respond receive the default option stated in the announcement."),
      ("Reconciliation", "Entitlements are reconciled against the custodian on the pay date. Breaks must be resolved within 3 business days.")]),
    ("ops_bcp_runbook", "operations", "sop", "confidential", "pdf",
     "Business Continuity Plan: Payments Runbook",
     [("Activation", "The business continuity plan is activated by the Head of Operations when the primary site is unavailable for more than 2 hours."),
      ("Recovery Objectives", "Payments systems have a recovery time objective (RTO) of 4 hours and a recovery point objective (RPO) of 15 minutes."),
      ("Alternate Site", "Critical staff relocate to the Charlotte recovery site. Call tree contact is managed by ops.bcp@meridian-example.com and the emergency line 555-0142.")]),

    # ------------------------------------------------------------------ wealth management
    ("wm_suitability_policy", "wealth_management", "financial_policy", "internal", "pdf",
     "Investment Suitability Policy",
     [("Purpose", "Advisors must ensure every recommendation is suitable for the client's investment profile. The policy is governed by FINRA Rule 2111 and Regulation Best Interest."),
      ("Investment Profile", "The investment profile includes age, income, net worth, investment objectives, risk tolerance, time horizon and liquidity needs. Profiles must be refreshed every 36 months or after a major life event."),
      ("Complex Products", "Leveraged ETFs, structured notes and private placements are complex products. They may only be recommended to clients with a risk tolerance of aggressive and require supervisory approval before the order is placed.")]),
    ("wm_client_onboarding_sop", "wealth_management", "sop", "internal", "docx",
     "SOP: Wealth Client Onboarding",
     [("Steps", "Onboarding has five steps: discovery meeting, KYC collection, investment profile, account opening and funding. This SOP requires the AML and KYC Policy and the Investment Suitability Policy."),
      ("Documents", "Collect a government ID, proof of address dated within 90 days, and a signed advisory agreement."),
      ("Service Level", "Accounts must be opened within 5 business days of receiving complete documentation. Example record: client Jane Doe, SSN 123-45-6789, email jane.doe@example.com, phone 555-0199.")]),
    ("wm_fee_schedule_kb", "wealth_management", "knowledge_base", "public", "html",
     "KB: Advisory Fee Schedule",
     [("Tiered Fees", "The annual advisory fee is 1.00 percent on the first 1 million USD, 0.75 percent on the next 4 million USD, and 0.50 percent on assets above 5 million USD."),
      ("Billing", "Fees are billed quarterly in advance, based on the account value on the last business day of the prior quarter."),
      ("Waivers", "Fee waivers above 25 percent require approval from the regional managing director.")]),

    # ------------------------------------------------------------------ asset servicing
    ("as_nav_calc_sop", "asset_servicing", "sop", "internal", "pdf",
     "SOP: Daily NAV Calculation",
     [("Timeline", "Fund accountants must strike the daily net asset value (NAV) by 6:00 PM Eastern Time. Pricing files from vendors arrive by 4:45 PM."),
      ("Price Tolerance", "Any security with a day-over-day price move greater than 5 percent is flagged as a price exception and must be validated against a second pricing source."),
      ("NAV Errors", "A NAV error of 0.50 percent or more for money market funds, or 1 percent or more for other funds, is material and requires shareholder reprocessing. Material NAV errors must be reported to the fund board.")]),
    ("as_custody_reconciliation_kb", "asset_servicing", "knowledge_base", "internal", "md",
     "KB: Custody Cash and Position Reconciliation",
     [("Frequency", "Cash and positions are reconciled daily against sub-custodians using the ReconOne platform."),
      ("Ageing", "Breaks older than 5 business days are reported to the Asset Servicing manager. Breaks older than 30 days are escalated to Operational Risk and require a written root cause."),
      ("Ownership", "Asset Servicing owns the reconciliation process; the Technology department owns the ReconOne platform.")]),
    ("as_securities_lending_policy", "asset_servicing", "financial_policy", "confidential", "docx",
     "Securities Lending Collateral Policy",
     [("Collateral", "Loans of US securities must be collateralized at 102 percent of market value; loans of non-US securities at 105 percent."),
      ("Marking to Market", "Collateral is marked to market daily. If coverage drops below the requirement, a collateral call is issued the same day."),
      ("Governance", "The Securities Lending Committee approves borrowers. The policy is governed by SEC Rule 15c3-3.")]),

    # ------------------------------------------------------------------ risk management
    ("risk_operational_risk_framework", "risk_management", "financial_policy", "internal", "pdf",
     "Operational Risk Management Framework",
     [("Three Lines", "Meridian Trust & Co. uses the three lines model. Business units are the first line, Risk Management and Compliance are the second line, and Internal Audit is the third line."),
      ("Incident Reporting", "Operational risk incidents with a loss of 10,000 USD or more must be reported in the Risk Event Register within 2 business days. Incidents with a loss of 1 million USD or more must be escalated to the Chief Risk Officer within 24 hours."),
      ("RCSA", "Each business unit completes a Risk and Control Self-Assessment (RCSA) annually.")]),
    ("risk_model_risk_policy", "risk_management", "financial_policy", "confidential", "html",
     "Model Risk Management Policy",
     [("Scope", "Applies to all quantitative models, including machine learning and generative AI models. The policy is governed by Federal Reserve guidance SR 11-7."),
      ("Validation", "High-tier models must be independently validated before production use and revalidated every 12 months. Medium-tier models are revalidated every 24 months."),
      ("Generative AI", "Generative AI use cases must document hallucination controls, human oversight and a monitoring plan before approval by the Model Risk Committee.")]),
    ("risk_vendor_risk_sop", "risk_management", "sop", "internal", "txt",
     "SOP: Third-Party Vendor Risk Assessment",
     [("Tiering", "Vendors are tiered as critical, high, moderate or low based on data access and service criticality."),
      ("Due Diligence", "Critical and high vendors require a SOC 2 Type II report, a security questionnaire and a financial review before contract signature."),
      ("Reassessment", "Critical vendors are reassessed annually; low-tier vendors every 3 years.")]),

    # ------------------------------------------------------------------ human resources
    ("hr_leave_policy", "human_resources", "financial_policy", "public", "md",
     "Paid Time Off and Leave Policy",
     [("Accrual", "Full-time employees accrue 20 days of paid time off per year. Up to 5 unused days can be carried over to the next year."),
      ("Mandatory Block Leave", "Employees in trading, payments and treasury roles must take a mandatory block leave of 10 consecutive business days each year as a fraud control."),
      ("Parental Leave", "Primary caregivers receive 16 weeks of paid parental leave; secondary caregivers receive 6 weeks.")]),
    ("hr_code_of_conduct", "human_resources", "financial_policy", "public", "pdf",
     "Code of Conduct",
     [("Speak Up", "Employees must report suspected misconduct to their manager, Compliance, or the anonymous Ethics Hotline. Retaliation against anyone who reports in good faith is prohibited."),
      ("Conflicts of Interest", "Outside business activities must be disclosed and approved before they begin."),
      ("Attestation", "All employees must attest to the Code of Conduct annually by January 31.")]),
    ("hr_compensation_bands", "human_resources", "financial_policy", "restricted", "docx",
     "Compensation Bands and Bonus Pool Guidelines",
     [("Bands", "Compensation bands are reviewed annually by the Compensation Committee. Band data is restricted to HR business partners and executives."),
      ("Bonus Pool", "The firm-wide bonus pool is funded at between 80 and 120 percent of target based on firm performance and risk-adjusted return."),
      ("Clawback", "Bonuses for material risk takers are subject to a 3-year clawback period.")]),

    # ------------------------------------------------------------------ technology
    ("tech_access_management_sop", "technology", "sop", "internal", "html",
     "SOP: User Access Management",
     [("Provisioning", "Access requests are submitted in the IAM portal and approved by the data owner. Privileged access additionally requires approval from the Information Security team."),
      ("Recertification", "Managers recertify their team's access every quarter. Privileged access is recertified monthly."),
      ("Leavers", "Access for leavers must be revoked within 24 hours of the termination date.")]),
    ("tech_genai_usage_policy", "technology", "financial_policy", "public", "md",
     "Generative AI Acceptable Use Policy",
     [("Approved Tools", "Only firm-approved generative AI tools, such as the Enterprise Knowledge Assistant, may be used with internal data."),
      ("Prohibited Data", "Client personal information, material non-public information and credentials must never be entered into external AI tools."),
      ("Human Review", "Outputs used in client communications must be reviewed by a human before sending. The policy is owned by the Technology department and governed by the Model Risk Management Policy.")]),
    ("tech_incident_response_kb", "technology", "knowledge_base", "confidential", "txt",
     "KB: Cybersecurity Incident Response",
     [("Severity", "Severity 1 incidents involve confirmed client data exposure or a ransomware event. Severity 1 incidents must be reported to the CISO within 1 hour."),
      ("Containment", "Isolate affected hosts from the network, preserve logs and disable compromised accounts."),
      ("Regulatory Notice", "If client data is affected, Legal decides on regulatory notification. Banking regulators must be notified within 36 hours of determining a notification incident.")]),
]

LEVEL_ORDER = ["public", "internal", "confidential", "restricted"]


def _plain(title: str, sections) -> str:
    out = [title, "=" * len(title), ""]
    for h, p in sections:
        out += [h, "-" * len(h), p, ""]
    return "\n".join(out)


def _markdown(title: str, sections) -> str:
    out = [f"# {title}", ""]
    for h, p in sections:
        out += [f"## {h}", "", p, ""]
    return "\n".join(out)


def _html(title: str, sections) -> str:
    body = "\n".join(f"<h2>{html.escape(h)}</h2>\n<p>{html.escape(p)}</p>" for h, p in sections)
    return (f"<!doctype html>\n<html><head><meta charset='utf-8'><title>{html.escape(title)}</title></head>\n"
            f"<body>\n<nav>Meridian Trust &amp; Co. intranet | Home | Policies</nav>\n"
            f"<h1>{html.escape(title)}</h1>\n{body}\n<footer>Internal use. Generated sample.</footer>\n</body></html>\n")


def _docx(path: Path, title: str, sections) -> None:
    from docx import Document

    doc = Document()
    doc.core_properties.title = title
    doc.add_heading(title, level=1)
    for h, p in sections:
        doc.add_heading(h, level=2)
        doc.add_paragraph(p)
    doc.save(path)


def _pdf(path: Path, title: str, sections) -> None:
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    story = [Paragraph(html.escape(title), styles["Title"]), Spacer(1, 12)]
    for h, p in sections:
        story += [Paragraph(html.escape(h), styles["Heading2"]), Paragraph(html.escape(p), styles["BodyText"]), Spacer(1, 8)]
    SimpleDocTemplate(str(path), pagesize=LETTER, title=title, invariant=1).build(story)


def main() -> None:
    rows = []
    for doc_id, bu, category, level, fmt, title, sections in DOCS:
        assert level in LEVEL_ORDER, level
        folder = RAW / bu
        folder.mkdir(parents=True, exist_ok=True)
        path = folder / f"{doc_id}.{fmt}"
        if fmt == "md":
            path.write_text(_markdown(title, sections), encoding="utf-8")
        elif fmt == "txt":
            path.write_text(_plain(title, sections), encoding="utf-8")
        elif fmt == "html":
            path.write_text(_html(title, sections), encoding="utf-8")
        elif fmt == "docx":
            _docx(path, title, sections)
        elif fmt == "pdf":
            _pdf(path, title, sections)
        else:
            raise ValueError(fmt)
        rows.append({
            "doc_id": doc_id, "title": title, "business_unit": bu, "category": category,
            "access_level": level, "file_format": fmt, "version": 1,
            "relative_path": path.relative_to(ROOT / "data" / "raw").as_posix(),
        })
    META.parent.mkdir(parents=True, exist_ok=True)
    with META.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)
    print(f"wrote {len(rows)} documents to {RAW}")


if __name__ == "__main__":
    main()
